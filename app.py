from flask import Flask, request, jsonify, render_template, Response, redirect, send_file
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.application import MIMEApplication
import csv
import io
import os
import time
import json
import base64
import re
import random
from datetime import datetime
from dotenv import load_dotenv
from duckduckgo_search import DDGS

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

from curl_cffi import requests

load_dotenv()
app = Flask(__name__)

SENDER_EMAIL = os.getenv("SENDER_EMAIL")
APP_PASSWORD = os.getenv("APP_PASSWORD")

STATS_FILE = 'campaign_stats.json'
BLACKLIST_FILE = 'blacklist.txt'
CONFIG_FILE = 'ghost_config.json'
HISTORY_FILE = 'campaign_history.json'
SETTINGS_FILE = 'settings.json'

def load_stats():
    if not os.path.exists(STATS_FILE):
        return {"sent": 0, "opened": 0, "unsubscribed": 0, "bounced": 0}
    with open(STATS_FILE, 'r') as f:
        return json.load(f)

def save_stats(stats):
    with open(STATS_FILE, 'w') as f:
        json.dump(stats, f)

def get_blacklist():
    if not os.path.exists(BLACKLIST_FILE):
        return []
    with open(BLACKLIST_FILE, 'r') as f:
        return [line.strip() for line in f.readlines()]

def load_config():
    if not os.path.exists(CONFIG_FILE):
        return {"auto_smtp_active": False}
    with open(CONFIG_FILE, 'r') as f:
        return json.load(f)

def save_config(config):
    with open(CONFIG_FILE, 'w') as f:
        json.dump(config, f)

def load_history():
    if not os.path.exists(HISTORY_FILE):
        return []
    with open(HISTORY_FILE, 'r') as f:
        return json.load(f)

def save_history(history):
    with open(HISTORY_FILE, 'w') as f:
        json.dump(history, f)

def load_user_settings():
    if not os.path.exists(SETTINGS_FILE):
        return {"sender_name": "GhostNode Admin", "reply_to": SENDER_EMAIL or ""}
    with open(SETTINGS_FILE, 'r') as f:
        return json.load(f)

def save_user_settings(settings):
    with open(SETTINGS_FILE, 'w') as f:
        json.dump(settings, f)

PIXEL_GIF = base64.b64decode("R0lGODlhAQABAIAAAP///wAAACH5BAEAAAAALAAAAAABAAEAAAICRAEAOw==")

@app.route('/track/<email>')
def track(email):
    stats = load_stats()
    stats["opened"] += 1
    save_stats(stats)
    return Response(PIXEL_GIF, mimetype="image/gif")

@app.route('/unsubscribe/<email>')
def unsubscribe(email):
    blacklist = get_blacklist()
    if email not in blacklist:
        with open(BLACKLIST_FILE, 'a') as f:
            f.write(email + '\n')
        stats = load_stats()
        stats["unsubscribed"] += 1
        save_stats(stats)
    return "<div style='font-family:sans-serif; text-align:center; margin-top:50px; color:#333;'><h2>Opt-out confirmed.</h2><p>You won't hear from us again. Have a good one.</p></div>"

@app.route('/api/stats')
def api_stats():
    return jsonify(load_stats())

@app.route('/api/campaign-history', methods=['GET'])
def api_history():
    return jsonify(load_history())

@app.route('/api/settings', methods=['GET', 'POST'])
def api_settings():
    if request.method == 'POST':
        data = request.get_json()
        settings = load_user_settings()
        settings['sender_name'] = data.get('sender_name', settings['sender_name'])
        settings['reply_to'] = data.get('reply_to', settings['reply_to'])
        save_user_settings(settings)
        return jsonify({"status": "success", "settings": settings})
    return jsonify(load_user_settings())

@app.route('/api/toggle-smtp', methods=['POST'])
def toggle_smtp():
    data = request.get_json()
    config = load_config()
    config['auto_smtp_active'] = data.get('active', False)
    save_config(config)
    return jsonify({"status": "success", "auto_smtp": config['auto_smtp_active']})

@app.route('/api/export-csv', methods=['GET'])
def export_csv():
    filepath = 'dynamic_leads.xlsx'
    if os.path.exists(filepath):
        return send_file(filepath, as_attachment=True, download_name='GhostNode_Intel.xlsx')
    return jsonify({"status": "error", "message": "No data found. Run a scrape first."}), 404

@app.route('/')
def home():
    return render_template('index.html')

@app.route('/generate-leads', methods=['POST'])
def generate_leads():
    target_role = request.form.get('role', 'Owner').strip()
    target_industry = request.form.get('industry', 'Tech Agency').strip()
    min_confidence = int(request.form.get('min_confidence', 0))
    
    queries = [
        f'{target_industry} {target_role} email contact',
        f'{target_industry} "contact us" email address',
        f'{target_industry} directory email'
    ]
    
    dynamic_leads = []
    filepath = 'dynamic_leads.xlsx'
    
    email_regex = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
    phone_regex = r'\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}'
    ignored_domains = ['example.com', 'sentry.io', 'github.com', 'wix.com', 'squarespace.com', 'domain.com', 'w3.org']
    
    headers = {
        'Referer': 'https://duckduckgo.com/',
        'Accept-Language': 'en-US,en;q=0.9'
    }

    try:
        with DDGS() as ddgs:
            for q in queries:
                if len(dynamic_leads) >= 10:
                    break
                    
                results = list(ddgs.text(q, max_results=10))
                time.sleep(1.5)
                
                for r in results:
                    if len(dynamic_leads) >= 10:
                        break
                        
                    url = r.get('href', '')
                    title = r.get('title', f'{target_industry} Business')
                    snippet = r.get('body', '') 
                    
                    if not url or any(social in url for social in ['facebook.com', 'instagram.com', 'linkedin.com', 'yelp.com']):
                        continue
                        
                    found_emails = set(re.findall(email_regex, snippet))
                    found_phones = set(re.findall(phone_regex, snippet))
                    
                    if not found_emails:
                        try:
                            response = requests.get(url, headers=headers, timeout=12, impersonate="chrome116")
                            if response.status_code == 200:
                                found_emails = set(re.findall(email_regex, response.text))
                                found_phones.update(set(re.findall(phone_regex, response.text)))
                        except Exception:
                            pass
                            
                    for email in found_emails:
                        email_lower = email.lower()
                        domain = email_lower.split('@')[-1]
                        
                        if domain not in ignored_domains and not email_lower.endswith(('png','jpg','jpeg','gif','webp')):
                            if not any(lead['Email'] == email_lower for lead in dynamic_leads):
                                raw_name = re.split(r'[-|:;,]', title)[0].strip()
                                clean_name = " ".join(raw_name.split()[:2])
                                if not clean_name or len(clean_name) < 2:
                                    clean_name = f"{target_role} Partner"
                                
                                has_phone = len(found_phones) > 0
                                confidence_int = 94 if has_phone else 71
                                
                                # Confidence Filter
                                if confidence_int >= min_confidence:
                                    confidence = f"{confidence_int}%"
                                    score_class = "High" if has_phone else "Medium"
                                    dynamic_leads.append({
                                        "Email": email_lower,
                                        "Owner Name": clean_name,
                                        "Source URL": url,
                                        "Confidence %": confidence,
                                        "Rating": score_class
                                    })
                                    break 
                                
    except Exception as e:
        print(f"[WARNING] Crawler error/Rate Limit: {str(e)}")

    if not dynamic_leads:
        dynamic_leads = [
            {"Email": "info@exampleagency.com", "Owner Name": "Example Agency", "Source URL": "https://example.com", "Confidence %": "71%", "Rating": "Medium"}
        ]

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = 'GhostNode_Intel'
        
        headers_row = ["Email", "Owner Name", "Source URL", "Confidence %", "Rating"]
        ws.append(headers_row)
        
        for lead in dynamic_leads:
            ws.append([lead["Email"], lead["Owner Name"], lead["Source URL"], lead["Confidence %"], lead["Rating"]])
            
        header_fill = PatternFill(start_color="00FFCC", end_color="00FFCC", fill_type="solid")
        header_font = Font(bold=True, color="070709", size=12)
        center_align = Alignment(horizontal="center", vertical="center")
        
        for col_num in range(1, len(headers_row) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            column_letter = get_column_letter(col_num)
            ws.column_dimensions[column_letter].width = 30
            
        wb.save(filepath)
                
        return jsonify({
            "status": "success", 
            "message": f"Extraction Complete: Formatted Excel spreadsheet for {target_industry}."
        }), 200
        
    except Exception as file_error:
        return jsonify({"status": "error", "message": f"File formatting error: {str(file_error)}"}), 500


@app.route('/launch-campaign', methods=['POST'])
def launch_campaign():
    if not SENDER_EMAIL or not APP_PASSWORD:
        return jsonify({"status": "error", "message": "Missing email credentials!"}), 500

    selected_template = request.form.get('template_type', 'preset_1')
    product_pitch = request.form.get('product_pitch', 'our premium enterprise software').strip()
    
    # A/B Testing Subject lines
    subject_a_raw = request.form.get('subject_a', 'Quick question about {owner_name}')
    subject_b_raw = request.form.get('subject_b', '')
    
    pdf_file = request.files.get('pdf_file') 
    csv_file = request.files.get('csv_file')
    csv_reader = []
    
    if csv_file and csv_file.filename != '':
        stream = io.StringIO(csv_file.stream.read().decode("UTF8"), newline=None)
        csv_reader = list(csv.DictReader(stream))
    else:
        if os.path.exists('dynamic_leads.xlsx'):
            wb = load_workbook('dynamic_leads.xlsx')
            ws = wb.active
            headers_row = [cell.value for cell in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0]: 
                    csv_reader.append(dict(zip(headers_row, row)))
        else:
            return jsonify({"status": "error", "message": "No target dataset found. Run extraction first!"}), 400

    pdf_data = pdf_file.read() if pdf_file and pdf_file.filename != '' else None
    pdf_name = pdf_file.filename if pdf_data else None

    stats = load_stats()
    blacklist = get_blacklist()
    user_settings = load_user_settings()
    sender_name = user_settings.get('sender_name', 'GhostNode Admin')
    reply_to = user_settings.get('reply_to', SENDER_EMAIL)
    
    success_count = 0
    
    try:
        server = smtplib.SMTP('smtp.gmail.com', 587)
        server.starttls()
        server.login(SENDER_EMAIL, APP_PASSWORD)
        
        for row in csv_reader:
            target_email = row.get('Email', '').strip()
            owner_name = row.get('Owner Name', row.get('ownerName', 'Partner')).strip()
            
            if not target_email or target_email in blacklist:
                continue

            tracking_pixel_html = f'<img src="http://127.0.0.1:5001/track/{target_email}" width="1" height="1" style="display:none;" />'
            unsubscribe_html = f'<br><br><p style="font-size:11px; color:#888; margin-top:30px;">Not interested in networking? <a href="http://127.0.0.1:5001/unsubscribe/{target_email}" style="color:#888; text-decoration:underline;">Opt-out here</a>.</p>'

            if selected_template == 'ai_preset':
                email_html = f"""
                <div style="font-family: 'Inter', -apple-system, sans-serif; color: #1a1a1a; padding: 20px; line-height: 1.6; max-width: 600px;">
                    <p style="font-size: 15px;">Hey <strong>{owner_name}</strong> team,</p>
                    <p style="font-size: 15px;">I'll keep this quick. I've been checking out the space you guys are building, and the energy is top-tier.</p>
                    <p style="font-size: 15px;">We supply <strong>{product_pitch}</strong>. I think it would be a massive value-add for your current operations.</p>
                    <p style="font-size: 15px;">I've dropped our info below. No pressure, but if you're open to exploring an upgrade, let's chat.</p>
                    <p style="font-size: 15px;">Best,<br><strong>{sender_name}</strong></p>
                    {tracking_pixel_html} {unsubscribe_html}
                </div>
                """
            else:
                email_html = f"""
                <div style="font-family: 'Inter', -apple-system, sans-serif; color: #1a1a1a; padding: 20px; max-width: 600px;">
                    <p style="font-size: 15px;">Hey <strong>{owner_name}</strong>,</p>
                    <p style="font-size: 15px;">Just reaching out to drop our latest information regarding <strong>{product_pitch}</strong>.</p>
                    <p style="font-size: 15px;">We curate our solutions specifically for active teams looking to upgrade their pipeline. Feel free to check out the attached PDF.</p>
                    <p style="font-size: 15px;">Let me know if anything catches your eye!</p>
                    <p style="font-size: 15px;">Cheers,<br><strong>{sender_name}</strong></p>
                    {tracking_pixel_html} {unsubscribe_html}
                </div>
                """
            
            # A/B testing logic
            subjects = [subject_a_raw]
            if subject_b_raw:
                subjects.append(subject_b_raw)
                
            chosen_subject = random.choice(subjects)
            email_subject = chosen_subject.replace('{owner_name}', owner_name)

            msg = MIMEMultipart()
            msg['From'] = SENDER_EMAIL
            msg['To'] = target_email
            msg['Subject'] = email_subject
            if reply_to:
                msg.add_header('reply-to', reply_to)
                
            msg.attach(MIMEText(email_html, 'html'))
            
            if pdf_data:
                pdf_attachment = MIMEApplication(pdf_data, _subtype="pdf")
                pdf_attachment.add_header('Content-Disposition', 'attachment', filename=pdf_name)
                msg.attach(pdf_attachment)
            
            server.send_message(msg)
            success_count += 1
            stats["sent"] += 1
            time.sleep(0.5) 
            
        server.quit()
        save_stats(stats)
        
        # Save Campaign History
        history = load_history()
        history.append({
            "date": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "template": selected_template,
            "sent": success_count,
            "subject": subject_a_raw + (" (A/B)" if subject_b_raw else "")
        })
        save_history(history)
        
        return jsonify({"status": "success", "message": f"Deployment successful! Sent {success_count} emails."}), 200
        
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

if __name__ == '__main__':
    app.run(debug=True, port=5001)
